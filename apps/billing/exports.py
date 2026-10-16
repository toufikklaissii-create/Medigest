"""Exports comptables des recettes (doctoresse uniquement). Aucune donnée clinique."""
import io
from decimal import Decimal

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.clinical.services import LETTERHEAD, doctor_display_name
from apps.common.pdf import render_pdf

from .reports import local_date

DH_FORMAT = '#,##0.00 "DH"'
HEADER_FILL = PatternFill("solid", fgColor="E3ECE7")
HEADER_FONT = Font(bold=True, color="2F3E46")
TOTAL_FONT = Font(bold=True)
TOTAL_BORDER = Border(top=Side(style="thin", color="6B9080"))


def export_filename(cabinet, period, extension):
    return f"recettes_{period.slug}_{slugify(cabinet.nom) or 'cabinet'}.{extension}"


def _table(sheet, headers, rows, money_columns=(), widths=()):
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    first_data_row = sheet.max_row + 1
    for row in rows:
        sheet.append(row)
    for column in money_columns:
        for (cell,) in sheet.iter_rows(min_row=first_data_row, min_col=column, max_col=column):
            cell.number_format = DH_FORMAT
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return first_data_row


def build_revenue_xlsx(cabinet, report):
    period = report["period"]
    workbook = Workbook()

    # ── Résumé ──
    summary = workbook.active
    summary.title = "Résumé"
    summary.append([f"Récapitulatif des recettes — {cabinet.nom}"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([f"Période : {period.label}"])
    summary.append([f"Édité le {timezone.localtime():%d/%m/%Y à %H:%M}"])
    summary.append([])
    first = _table(
        summary, ["Indicateur", "Valeur"],
        [
            ["Total encaissé", report["total"]],
            ["Nombre d'encaissements", report["count"]],
            ["Patientes distinctes", report["patients"]],
            ["Panier moyen", report["average"]],
            ["Impayés en cours (reste dû)", report["outstanding"]],
            ["Factures annulées", report["cancelled"]],
        ],
        widths=(34, 18),
    )
    for offset in (0, 3, 4):
        summary.cell(row=first + offset, column=2).number_format = DH_FORMAT
    summary.append([])
    first = _table(summary, ["Mode de paiement", "Montant"], [[m["label"], m["amount"]] for m in report["by_method"]], money_columns=(2,))
    last = summary.max_row
    _sum_row(summary, "Total", first, last, money=(2,))
    summary.freeze_panes = "A6"

    # ── Encaissements ──
    sheet = workbook.create_sheet("Encaissements")
    rows = [
        [
            local_date(p.paid_at),
            timezone.localtime(p.paid_at).strftime("%H:%M"),
            p.invoice.invoice_number,
            f"{p.invoice.patient.last_name} {p.invoice.patient.first_name}",
            p.acts_label,
            p.get_method_display(),
            p.reference,
            p.amount,
            (p.received_by.get_full_name() or p.received_by.username) if p.received_by else "",
        ]
        for p in report["payments"]
    ]
    first = _table(
        sheet,
        ["Date", "Heure", "N° facture", "Patiente", "Actes", "Mode", "Référence", "Montant", "Encaissé par"],
        rows, money_columns=(8,), widths=(12, 8, 16, 26, 36, 12, 16, 14, 20),
    )
    for (cell,) in sheet.iter_rows(min_row=first, min_col=1, max_col=1):
        cell.number_format = "DD/MM/YYYY"
    _sum_row(sheet, "Total", first, sheet.max_row, money=(8,))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(first - 1 + len(rows), 1)}"

    # ── Par acte ──
    sheet = workbook.create_sheet("Par acte")
    first = _table(
        sheet, ["Acte", "Quantité", "Montant encaissé"],
        [[row["name"], row["quantity"], row["amount"]] for row in report["by_act"]],
        money_columns=(3,), widths=(40, 12, 18),
    )
    _sum_row(sheet, "Total", first, sheet.max_row, money=(3,), counts=(2,))
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sum_row(sheet, label, first_row, last_row, money=(), counts=()):
    """Ligne de total, en valeurs (lisible même par un lecteur qui ne recalcule pas les formules)."""
    row = sheet.max_row + 1
    sheet.cell(row=row, column=1, value=label)
    for column in (*money, *counts):
        values = [sheet.cell(row=r, column=column).value or 0 for r in range(first_row, last_row + 1)]
        cell = sheet.cell(row=row, column=column, value=sum(values, Decimal("0") if column in money else 0))
        if column in money:
            cell.number_format = DH_FORMAT
    for cell in sheet[row]:
        cell.font, cell.border = TOTAL_FONT, TOTAL_BORDER
    return row


def build_revenue_pdf(cabinet, report, doctor):
    return render_pdf(
        "pdf/revenue_summary.html",
        {
            **report,
            "cabinet": cabinet,
            "doc_date": timezone.localtime(),
            "doctor_name": doctor_display_name(doctor),
        },
        background=LETTERHEAD,
    )
