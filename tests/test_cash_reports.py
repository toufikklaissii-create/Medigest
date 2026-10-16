import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal
from unittest import mock
from zoneinfo import ZoneInfo

from django.contrib.auth.models import User
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import Cabinet, Profile
from apps.billing.models import CashClosing, Invoice, InvoiceItem, MedicalAct, Payment
from apps.billing.reports import cash_journal, make_period, period_payments, revenue_report
from apps.patients.models import Patient

CASA = ZoneInfo("Africa/Casablanca")


def local(*args):
    return datetime(*args, tzinfo=CASA)


@override_settings(ALLOWED_HOSTS=["testserver"])
class CashReportsBase(TestCase):
    def setUp(self):
        self.cabinet = Cabinet.objects.create(nom="Cabinet Dr ERRMILI", slug="cabinet-recettes")
        self.doctor = User.objects.create_user(username="dr_rec", password="x", first_name="Khadija", last_name="Errmili")
        Profile.objects.update_or_create(user=self.doctor, defaults={"role": Profile.Role.DOCTOR, "cabinet": self.cabinet})
        self.assistant = User.objects.create_user(username="as_rec", password="x")
        Profile.objects.update_or_create(user=self.assistant, defaults={"role": Profile.Role.ASSISTANT, "cabinet": self.cabinet})
        self.patient = Patient.objects.create(cabinet=self.cabinet, first_name="Nadia", last_name="Alaoui")
        self.other_patient = Patient.objects.create(cabinet=self.cabinet, first_name="Salma", last_name="Bennani")
        self.consult = MedicalAct.objects.create(cabinet=self.cabinet, code="CONS", name="Consultation", price=Decimal("300"))
        self.echo = MedicalAct.objects.create(cabinet=self.cabinet, code="ECHO", name="Échographie", price=Decimal("200"))
        self.doctor_client = Client()
        self.doctor_client.force_login(self.doctor)
        self.assistant_client = Client()
        self.assistant_client.force_login(self.assistant)

    def invoice(self, acts=None, patient=None, discount=Decimal("0")):
        acts = acts or [self.consult]
        invoice = Invoice.objects.create(
            cabinet=self.cabinet, patient=patient or self.patient,
            total_amount=sum((a.price for a in acts), Decimal("0")) - discount, discount=discount,
        )
        for act in acts:
            InvoiceItem.objects.create(invoice=invoice, act=act, price=act.price)
        return invoice

    def pay(self, invoice, amount, paid_at, method=Payment.Method.CASH, reference=""):
        payment = Payment.objects.create(
            invoice=invoice, amount=Decimal(amount), method=method, reference=reference,
            paid_at=paid_at, received_by=self.assistant,
        )
        invoice.refresh_payment_status()
        return payment


class RevenueTotalsTest(CashReportsBase):
    def test_midnight_and_month_boundaries_use_casablanca_time(self):
        # 23:30 le 31/10 à Casablanca = 22:30 UTC ; 00:15 le 1/11 = 31/10 23:15 UTC.
        self.pay(self.invoice(), "300", local(2026, 10, 31, 23, 30))
        self.pay(self.invoice(patient=self.other_patient), "200", local(2026, 11, 1, 0, 15))

        october = revenue_report(self.cabinet, make_period("mois", date(2026, 10, 5)))
        november = revenue_report(self.cabinet, make_period("mois", date(2026, 11, 20)))
        self.assertEqual(october["total"], Decimal("300"))
        self.assertEqual(november["total"], Decimal("200"))
        self.assertEqual(cash_journal(self.cabinet, date(2026, 10, 31))["total"], Decimal("300"))
        self.assertEqual(cash_journal(self.cabinet, date(2026, 11, 1))["total"], Decimal("200"))
        # Semaine du lundi 26/10 au dimanche 1/11 : les deux paiements.
        week = revenue_report(self.cabinet, make_period("semaine", date(2026, 10, 28)))
        self.assertEqual((week["period"].start, week["period"].end), (date(2026, 10, 26), date(2026, 11, 1)))
        self.assertEqual(week["total"], Decimal("500"))
        self.assertEqual(week["patients"], 2)
        by_day = {row["date"]: row["amount"] for row in week["by_day"]}
        self.assertEqual(by_day[date(2026, 10, 31)], Decimal("300"))
        self.assertEqual(by_day[date(2026, 11, 1)], Decimal("200"))

    def test_partial_payments_count_at_their_own_date(self):
        invoice = self.invoice([self.consult, self.echo])  # 500 DH
        self.pay(invoice, "200", local(2026, 10, 14, 10))
        self.pay(invoice, "300", local(2026, 10, 15, 11), method=Payment.Method.CARD)
        self.assertEqual(cash_journal(self.cabinet, date(2026, 10, 14))["total"], Decimal("200"))
        day15 = cash_journal(self.cabinet, date(2026, 10, 15))
        self.assertEqual(day15["total"], Decimal("300"))
        self.assertEqual({r["method"]: r["amount"] for r in day15["by_method"]}["card"], Decimal("300"))
        month = revenue_report(self.cabinet, make_period("mois", date(2026, 10, 1)))
        self.assertEqual(month["total"], Decimal("500"))
        self.assertEqual(month["count"], 2)
        self.assertEqual(month["average"], Decimal("250.00"))
        # Acte compté une seule fois malgré deux paiements ; montants = total encaissé.
        acts = {row["name"]: row for row in month["by_act"]}
        self.assertEqual(acts["Consultation"]["quantity"], 1)
        self.assertEqual(acts["Consultation"]["amount"], Decimal("300.00"))
        self.assertEqual(acts["Échographie"]["amount"], Decimal("200.00"))

    def test_act_breakdown_with_discount_matches_total(self):
        invoice = self.invoice([self.consult, self.echo], discount=Decimal("100"))  # 400 DH
        self.pay(invoice, "133.33", local(2026, 10, 10, 9))
        self.pay(invoice, "266.67", local(2026, 10, 12, 9))
        report = revenue_report(self.cabinet, make_period("mois", date(2026, 10, 1)))
        self.assertEqual(sum(row["amount"] for row in report["by_act"]), Decimal("400.00"))
        self.assertEqual(report["by_act"][0]["name"], "Consultation")
        self.assertEqual(report["by_act"][0]["amount"], Decimal("240.00"))

    def test_cancelled_invoices_are_excluded(self):
        kept = self.invoice()
        self.pay(kept, "300", local(2026, 10, 10, 9))
        cancelled = self.invoice(patient=self.other_patient)
        self.pay(cancelled, "300", local(2026, 10, 10, 10))
        Invoice.objects.filter(pk=cancelled.pk).update(
            status=Invoice.Status.CANCELLED, cancelled_at=local(2026, 10, 11, 9), cancel_reason="Erreur"
        )
        report = revenue_report(self.cabinet, make_period("mois", date(2026, 10, 1)))
        self.assertEqual(report["total"], Decimal("300"))
        self.assertEqual(report["count"], 1)
        self.assertEqual(report["cancelled"], 1)
        self.assertEqual(period_payments(self.cabinet, date(2026, 10, 10), date(2026, 10, 10)).count(), 1)

    def test_outstanding_and_previous_period_comparison(self):
        self.pay(self.invoice(), "300", local(2026, 9, 20, 9))
        self.pay(self.invoice(), "200", local(2026, 10, 3, 9))
        self.pay(self.invoice(patient=self.other_patient), "250", local(2026, 10, 4, 9))
        self.invoice([self.echo])  # 200 DH jamais encaissés
        report = revenue_report(self.cabinet, make_period("mois", date(2026, 10, 1)))
        self.assertEqual(report["previous"]["total"], Decimal("300"))
        self.assertAlmostEqual(report["variations"]["total"], 50.0)
        self.assertEqual(report["outstanding"], Decimal("350"))  # 100 + 50 restants + 200 non payés
        self.assertEqual(make_period("mois", date(2026, 3, 15)).previous().start, date(2026, 2, 1))


class CashClosingTest(CashReportsBase):
    def test_journal_lists_payments_and_due_invoices(self):
        today = timezone.localdate()
        invoice = self.invoice([self.consult, self.echo])
        self.pay(invoice, "200", timezone.now(), method=Payment.Method.CHECK, reference="CHQ-881")
        response = self.assistant_client.get(reverse("billing:cash_register"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, invoice.invoice_number)
        self.assertContains(response, "CHQ-881")
        self.assertContains(response, "Encore dues ce jour")
        journal = cash_journal(self.cabinet, today)
        self.assertEqual(journal["patients_count"], 1)
        self.assertEqual(journal["due_total"], Decimal("300"))
        self.assertEqual(journal["payments"][0].acts_label, "Consultation, Échographie")

    def test_close_day_records_difference_and_locks(self):
        self.pay(self.invoice(), "300", timezone.now())
        self.pay(self.invoice(patient=self.other_patient), "200", timezone.now(), method=Payment.Method.CARD)
        response = self.assistant_client.post(
            reverse("billing:cash_close"),
            {"date": timezone.localdate().isoformat(), "counted_cash": "280,00", "note": "Monnaie rendue"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        closing = CashClosing.objects.get(cabinet=self.cabinet)
        self.assertEqual(closing.expected_cash, Decimal("300"))  # espèces seulement
        self.assertEqual(closing.counted_cash, Decimal("280"))
        self.assertEqual(closing.difference, Decimal("-20"))
        self.assertEqual(closing.closed_by, self.assistant)
        self.assertIn("écart", json.loads(response.headers["HX-Trigger"])["toast"]["message"])
        self.assertContains(response, "Journée clôturée")

        again = self.assistant_client.post(
            reverse("billing:cash_close"), {"date": timezone.localdate().isoformat(), "counted_cash": "300"}
        )
        self.assertContains(again, "déjà clôturée")
        self.assertEqual(CashClosing.objects.count(), 1)

    def test_future_day_and_missing_amount_refused(self):
        tomorrow = (timezone.localdate() + timedelta(days=1)).isoformat()
        self.assistant_client.post(reverse("billing:cash_close"), {"date": tomorrow, "counted_cash": "0"})
        self.assistant_client.post(reverse("billing:cash_close"), {"date": timezone.localdate().isoformat(), "counted_cash": ""})
        self.assertFalse(CashClosing.objects.exists())

    def test_payment_after_closing_is_allowed_but_flagged(self):
        early = self.pay(self.invoice(), "300", timezone.now() - timedelta(minutes=10))
        Payment.objects.filter(pk=early.pk).update(created_at=timezone.now() - timedelta(minutes=10))
        closed_at = timezone.now() - timedelta(minutes=5)
        CashClosing.objects.create(
            cabinet=self.cabinet, date=timezone.localdate(), expected_cash=Decimal("300"),
            counted_cash=Decimal("300"), difference=Decimal("0"), closed_by=self.assistant, closed_at=closed_at,
        )
        late = self.invoice(patient=self.other_patient)
        response = self.assistant_client.post(
            reverse("billing:pay_invoice", args=[late.pk]), {"amount": "300", "payment_method": "cash"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("clôturée", json.loads(response.headers["HX-Trigger"])["toast"]["message"])
        late.refresh_from_db()
        self.assertEqual(late.status, Invoice.Status.PAID)
        journal = cash_journal(self.cabinet, timezone.localdate())
        self.assertEqual([p.after_closing for p in journal["payments"]], [False, True])
        self.assertEqual(journal["cash_since_closing"], Decimal("300"))

    def test_closing_history_visible_to_doctor_only(self):
        CashClosing.objects.create(
            cabinet=self.cabinet, date=date(2026, 10, 1), expected_cash=Decimal("100"),
            counted_cash=Decimal("90"), difference=Decimal("-10"), note="Billet manquant",
        )
        self.assertContains(self.doctor_client.get(reverse("billing:cash_register")), "Billet manquant")
        self.assertNotContains(self.assistant_client.get(reverse("billing:cash_register")), "Historique des clôtures")


class ReportsAccessAndExportTest(CashReportsBase):
    def test_assistant_forbidden_on_reports_and_exports(self):
        for name in ("billing:reports", "billing:export_xlsx", "billing:export_pdf"):
            with self.subTest(name=name):
                self.assertEqual(self.assistant_client.get(reverse(name)).status_code, 403)

    def test_reports_page_for_doctor(self):
        self.pay(self.invoice(), "300", local(2026, 10, 10, 9))
        response = self.doctor_client.get(reverse("billing:reports"), {"periode": "mois", "date": "2026-10-01"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "300,00")
        partial = self.doctor_client.get(
            reverse("billing:reports"), {"periode": "libre", "du": "2026-10-01", "au": "2026-10-15"}, HTTP_HX_REQUEST="true"
        )
        self.assertContains(partial, 'id="report-body"')

    def test_xlsx_export_sheets_and_sums(self):
        invoice = self.invoice([self.consult, self.echo])
        self.pay(invoice, "200", local(2026, 10, 14, 10))
        self.pay(invoice, "300", local(2026, 10, 15, 11), method=Payment.Method.TRANSFER, reference="VIR-42")
        self.pay(self.invoice(patient=self.other_patient), "300", local(2026, 10, 20, 9), method=Payment.Method.CARD)
        self.pay(self.invoice(), "999", local(2026, 11, 2, 9))  # hors période

        response = self.doctor_client.get(reverse("billing:export_xlsx"), {"periode": "mois", "date": "2026-10-01"})
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="recettes_2026-10_cabinet-dr-errmili.xlsx"', response["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Résumé", "Encaissements", "Par acte"])

        sheet = workbook["Encaissements"]
        headers = [c.value for c in sheet[1]]
        self.assertEqual(headers, ["Date", "Heure", "N° facture", "Patiente", "Actes", "Mode", "Référence", "Montant", "Encaissé par"])
        self.assertEqual(sheet.freeze_panes, "A2")
        amounts = [row[7] for row in sheet.iter_rows(min_row=2, max_row=4, values_only=True)]
        self.assertTrue(all(isinstance(v, (int, float)) for v in amounts))
        self.assertEqual(sum(amounts), 800)
        total_row = list(sheet.iter_rows(min_row=5, max_row=5, values_only=True))[0]
        self.assertEqual((total_row[0], total_row[7]), ("Total", 800))
        self.assertIn("VIR-42", [row[6] for row in sheet.iter_rows(min_row=2, values_only=True)])
        self.assertEqual(sheet["H2"].number_format, '#,##0.00 "DH"')

        acts = workbook["Par acte"]
        rows = list(acts.iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[-1][0], "Total")
        self.assertEqual(rows[-1][2], 800)
        self.assertEqual(sum(r[2] for r in rows[:-1]), 800)

        summary = workbook["Résumé"]
        values = {row[0]: row[1] for row in summary.iter_rows(values_only=True) if row and row[0]}
        self.assertEqual(values["Total encaissé"], 800)
        self.assertEqual(values["Nombre d'encaissements"], 3)
        self.assertEqual(values["Carte"], 300)

        all_text = " ".join(
            str(v) for ws in workbook for row in ws.iter_rows(values_only=True) for v in row if v is not None
        ).lower()
        self.assertNotIn("diagnostic", all_text)

    def test_pdf_summary_export(self):
        self.pay(self.invoice(), "300", local(2026, 10, 10, 9))
        with mock.patch("apps.billing.exports.render_pdf", return_value=b"%PDF-fake") as render:
            response = self.doctor_client.get(reverse("billing:export_pdf"), {"periode": "semaine", "date": "2026-10-10"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("recettes_2026-S41_cabinet-dr-errmili.pdf", response["Content-Disposition"])
        context = render.call_args.args[1]
        self.assertEqual(context["total"], Decimal("300"))

    def test_pdf_summary_renders_for_real(self):
        self.pay(self.invoice(), "300", local(2026, 10, 10, 9))
        response = self.doctor_client.get(reverse("billing:export_pdf"), {"periode": "mois", "date": "2026-10-01"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
