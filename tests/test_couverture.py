"""Couverture médicale (CNOPS / CNSS / sans couverture) : saisie, migration des fiches
existantes, liste « à vérifier » et séparation des recettes pour le comptable."""
import io
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.contrib.auth.models import User
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test import Client, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Cabinet, Profile
from apps.billing.models import Invoice, InvoiceItem, MedicalAct, Payment
from apps.billing.reports import make_period, revenue_report
from apps.patients.couverture import AUCUNE, AUTRE, CNOPS, CNSS, deviner_couverture
from apps.patients.models import Patient

CASA = ZoneInfo("Africa/Casablanca")


class DevinerCouvertureTest(TestCase):
    def test_cas_certains(self):
        cases = {
            "CNSS": CNSS, "cnss": CNSS, "C.N.S.S": CNSS, " Cnss AMO ": CNSS, "CNSS n° 123456": CNSS,
            "CNOPS": CNOPS, "cnops": CNOPS, "C.N.O.P.S.": CNOPS, "Cnops - MGPAP": CNOPS,
            "aucune": AUCUNE, "Aucun": AUCUNE, "sans": AUCUNE, "Sans couverture": AUCUNE, "néant": AUCUNE,
            "-": "", "": "", "   ": "",
        }
        for texte, attendu in cases.items():
            with self.subTest(texte=texte):
                self.assertEqual(deviner_couverture(texte), attendu)

    def test_cas_douteux_restent_a_verifier(self):
        for texte in ("AMO", "mutuelle", "Saham assurance", "FAR", "CNSS ou CNOPS ?", "MGPAP", "privée"):
            with self.subTest(texte=texte):
                self.assertEqual(deviner_couverture(texte), "")


class MigrationCouvertureTest(TransactionTestCase):
    """La migration range les fiches existantes sans jamais modifier le texte saisi."""

    migrate_from = [("patients", "0005_alter_patient_patient_number_and_more")]
    migrate_to = [("patients", "0006_couverture_medicale")]

    def test_fiches_existantes_classees(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_from)
        old_apps = executor.loader.project_state(self.migrate_from).apps
        OldPatient = old_apps.get_model("patients", "Patient")
        textes = {"CNSS": CNSS, "c.n.o.p.s": CNOPS, "aucune": AUCUNE, "AMO": "", "Mutuelle X": "", "": ""}
        ids = {}
        for index, texte in enumerate(textes):
            ids[texte] = OldPatient.objects.create(
                first_name=f"Fictive{index}", last_name="Test", patient_number=f"T-{index}", organisme_social=texte,
            ).pk

        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(self.migrate_to)
        new_apps = executor.loader.project_state(self.migrate_to).apps
        NewPatient = new_apps.get_model("patients", "Patient")
        for texte, attendu in textes.items():
            patient = NewPatient.objects.get(pk=ids[texte])
            self.assertEqual(patient.couverture, attendu, texte)
            self.assertEqual(patient.organisme_social, texte)  # texte d'origine conservé

        # Remet la base dans l'état final pour les tests suivants.
        executor = MigrationExecutor(connection)
        executor.loader.build_graph()
        executor.migrate(executor.loader.graph.leaf_nodes())


@override_settings(ALLOWED_HOSTS=["testserver"])
class CouvertureSaisieTest(TestCase):
    def setUp(self):
        self.cabinet = Cabinet.objects.create(nom="Cabinet Test", slug="cabinet-couv")
        self.assistant = User.objects.create_user(username="as_couv", password="x")
        Profile.objects.update_or_create(
            user=self.assistant, defaults={"role": Profile.Role.ASSISTANT, "cabinet": self.cabinet}
        )
        self.client = Client()
        self.client.force_login(self.assistant)

    def test_creation_avec_couverture(self):
        self.client.post(reverse("patients:create_patient"), {
            "first_name": "Fictive", "last_name": "Une", "couverture": CNOPS, "organisme_social": "MGPAP",
        })
        patient = Patient.objects.get(first_name="Fictive")
        self.assertEqual((patient.couverture, patient.organisme_social), (CNOPS, "MGPAP"))

    def test_valeur_inconnue_ignoree(self):
        self.client.post(reverse("patients:create_patient"), {
            "first_name": "Fictive", "last_name": "Deux", "couverture": "pirate",
        })
        self.assertEqual(Patient.objects.get(first_name="Fictive").couverture, "")

    def test_modification(self):
        patient = Patient.objects.create(cabinet=self.cabinet, first_name="Fictive", last_name="Trois")
        self.client.post(reverse("patients:edit_patient", args=[patient.pk]), {
            "first_name": "Fictive", "last_name": "Trois", "couverture": CNSS,
        })
        patient.refresh_from_db()
        self.assertEqual(patient.couverture, CNSS)

    def test_formulaire_propose_les_choix(self):
        html = self.client.get(reverse("patients:create_patient")).content.decode()
        for label in ("CNOPS", "CNSS", "Sans couverture", "Autre (mutuelle, FAR, assurance…)"):
            self.assertIn(label, html)
        self.assertIn('name="couverture"', html)

    def test_liste_a_verifier(self):
        Patient.objects.create(cabinet=self.cabinet, first_name="Avecc", last_name="Couverture", couverture=CNSS)
        Patient.objects.create(cabinet=self.cabinet, first_name="Sansinfo", last_name="Fiche", organisme_social="AMO")
        other = Cabinet.objects.create(nom="Autre", slug="autre-couv")
        Patient.objects.create(cabinet=other, first_name="Autrecabinet", last_name="X")
        html = self.client.get(reverse("patients:index"), {"view": "a_verifier"}).content.decode()
        self.assertIn("Sansinfo", html)
        self.assertNotIn("Avecc", html)
        self.assertNotIn("Autrecabinet", html)
        self.assertIn("Couverture à vérifier", html)
        self.assertIn("À vérifier", html)


@override_settings(ALLOWED_HOSTS=["testserver"])
class RecettesParCouvertureTest(TestCase):
    def setUp(self):
        self.cabinet = Cabinet.objects.create(nom="Cabinet Dr ERRMILI", slug="cabinet-couv-rec")
        self.doctor = User.objects.create_user(username="dr_couv", password="x")
        Profile.objects.update_or_create(user=self.doctor, defaults={"role": Profile.Role.DOCTOR, "cabinet": self.cabinet})
        self.act = MedicalAct.objects.create(cabinet=self.cabinet, code="CONS", name="Consultation", price=Decimal("300"))
        self.client = Client()
        self.client.force_login(self.doctor)
        self.patients = {
            code: Patient.objects.create(cabinet=self.cabinet, first_name=f"Fictive{index}", last_name="Couv",
                                         couverture=code)
            for index, code in enumerate((CNOPS, CNSS, AUCUNE, AUTRE, ""))
        }
        amounts = {CNOPS: "300", CNSS: "250", AUCUNE: "200", AUTRE: "150", "": "100"}
        for code, amount in amounts.items():
            invoice = Invoice.objects.create(cabinet=self.cabinet, patient=self.patients[code],
                                             total_amount=Decimal(amount))
            InvoiceItem.objects.create(invoice=invoice, act=self.act, price=Decimal(amount))
            Payment.objects.create(invoice=invoice, amount=Decimal(amount), method=Payment.Method.CASH,
                                   paid_at=datetime(2026, 10, 10, 10, tzinfo=CASA))
        # La CNSS a une 2e recette, pour vérifier les sommes.
        invoice = Invoice.objects.create(cabinet=self.cabinet, patient=self.patients[CNSS], total_amount=Decimal("50"))
        Payment.objects.create(invoice=invoice, amount=Decimal("50"), method=Payment.Method.CARD,
                               paid_at=datetime(2026, 10, 12, 10, tzinfo=CASA))

    def test_rapport_par_couverture(self):
        report = revenue_report(self.cabinet, make_period("mois", datetime(2026, 10, 1).date()))
        rows = {row["code"]: row for row in report["by_coverage"]}
        self.assertEqual([row["code"] for row in report["by_coverage"]], [CNOPS, CNSS, AUCUNE, AUTRE, ""])
        self.assertEqual(rows[CNSS]["amount"], Decimal("300"))
        self.assertEqual((rows[CNSS]["count"], rows[CNSS]["patients"]), (2, 1))
        self.assertEqual(rows[CNOPS]["amount"], Decimal("300"))
        self.assertEqual(rows[""]["label"], "Non renseignée")
        self.assertEqual(sum(row["amount"] for row in rows.values()), report["total"])

    def test_groupes_principaux_toujours_presents(self):
        Payment.objects.all().delete()
        report = revenue_report(self.cabinet, make_period("mois", datetime(2026, 10, 1).date()))
        self.assertEqual([row["code"] for row in report["by_coverage"]], [CNOPS, CNSS, AUCUNE])

    def test_page_rapports(self):
        html = self.client.get(reverse("billing:reports"), {"periode": "mois", "date": "2026-10-01"}).content.decode()
        self.assertIn("Par couverture", html)
        self.assertIn("à compléter", html)

    def test_export_excel_un_onglet_par_couverture(self):
        response = self.client.get(reverse("billing:export_xlsx"), {"periode": "mois", "date": "2026-10-01"})
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, [
            "Résumé", "Encaissements", "CNOPS", "CNSS", "Sans couverture", "Autre", "Non renseignée", "Par acte",
        ])
        cnss = list(workbook["CNSS"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(cnss), 3)  # 2 encaissements + total
        self.assertEqual(cnss[-1][0], "Total")
        self.assertEqual(cnss[-1][9], 300)
        self.assertTrue(all(row[4] == "CNSS" for row in cnss[:-1]))
        summary = {row[0]: row[1] for row in workbook["Résumé"].iter_rows(values_only=True) if row and row[0]}
        self.assertEqual(summary["CNOPS"], 300)
        self.assertEqual(summary["Sans couverture"], 200)

    def test_pdf_comptable(self):
        response = self.client.get(reverse("billing:export_pdf"), {"periode": "mois", "date": "2026-10-01"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
